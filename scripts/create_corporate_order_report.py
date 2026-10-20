from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Personal Projects\renivet-marketplace")
OUT = ROOT / "outputs" / "Renivet_Corporate_Order_Current_State_Report_With_Buyer_And_Admin_UI.docx"
IMG = ROOT / "tmp" / "corp-docs-png"
UI_IMG = ROOT / "tmp" / "corporate-ui-screenshots"
GREEN = "3F6045"; INK = "172033"; MUTED = "5F6B7A"; LINE = "D9DFD5"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def borders(table, color=LINE, size="6"):
    tblPr = table._tbl.tblPr
    b = tblPr.first_child_found_in("w:tblBorders")
    if b is None:
        b = OxmlElement("w:tblBorders"); tblPr.append(b)
    for edge in ("top","left","bottom","right","insideH","insideV"):
        el = b.find(qn("w:" + edge))
        if el is None:
            el = OxmlElement("w:" + edge); b.append(el)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0"); el.set(qn("w:color"), color)

def set_cell_text(cell, text, bold=False, color="172033", size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
    r = p.add_run(str(text)); r.bold = bold; r.font.name = "Aptos"
    r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, True, "FFFFFF")
        shade(t.rows[0].cells[i], GREEN)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
            if ri % 2 == 1: shade(cells[i], "F7F9F5")
    hdr = OxmlElement("w:tblHeader"); t.rows[0]._tr.get_or_add_trPr().append(hdr)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p

def para(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12; p.add_run(text)
    return p

def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(3)
        p.add_run(item)

def caption(doc, text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text); r.italic = True; r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(MUTED)

def picture(doc, filename, text, width=5.55):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(IMG / filename), width=Inches(width))
    caption(doc, text)

def ui_picture(doc, filename, text, width=6.35):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(UI_IMG / filename), width=Inches(width))
    caption(doc, text)

doc = Document()
sec = doc.sections[0]; sec.top_margin = Inches(0.65); sec.bottom_margin = Inches(0.62)
sec.left_margin = Inches(0.72); sec.right_margin = Inches(0.72)
styles = doc.styles
styles["Normal"].font.name = "Aptos"; styles["Normal"].font.size = Pt(9.6); styles["Normal"].font.color.rgb = RGBColor.from_string(INK)
styles["Title"].font.name = "Aptos Display"; styles["Title"].font.size = Pt(30); styles["Title"].font.bold = True
for s, size in [("Heading 1",18),("Heading 2",12.5),("Heading 3",10.5)]:
    styles[s].font.name = "Aptos Display"; styles[s].font.size = Pt(size); styles[s].font.bold = True
footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("Renivet Corporate Order Current State | 06 October 2026").font.size = Pt(8)

p = doc.add_paragraph(style="Title"); p.add_run("Renivet Corporate Order Current State")
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(18)
r = p.add_run("Current UI, features, order flow, documents, and architecture"); r.font.size = Pt(15); r.font.color.rgb = RGBColor.from_string(GREEN)
para(doc, "This document records the current corporate-order implementation in the repository: the buyer journeys, order states, admin and brand workspaces, payment flow, document chain, and system architecture.")
add_table(doc, ["Report detail","Value"], [
    ["Repository reviewed", r"C:\Personal Projects\renivet-marketplace"],
    ["Review date", "06 October 2026"],
    ["Scope", "Corporate ordering, RFQ, payment, fulfilment, finance, invoices, and operational documents"],
    ["Evidence used", "Current React/Next.js routes and components plus rendered PDF previews stored in the repository"],
])
para(doc, "Current implementation summary: the system contains a legacy self-service corporate order journey and a newer managed procurement journey. Both connect to corporate order records, payment records, operational status, and downloadable corporate documents.")
doc.add_page_break()

heading(doc, "1 Current state at a glance")
para(doc, "The current system contains two overlapping buyer experiences. The legacy self-service flow is a five-step configurable garment ordering journey. The newer corporate procurement surface is a dashboard-led RFQ, quote, purchase order, payment, fulfilment, and finance journey. Both are present in the repository and are linked from the corporate landing surface.")
add_table(doc, ["Area","Implemented today","Current evidence"], [
    ["Buyer entry", "Corporate procurement landing page, corporate orders page, profile corporate dashboard, legacy self-service route", "Both the newer procurement links and legacy self-service link are present."],
    ["Order capture", "Company/contact, delivery, quantity, product configuration, colours, branding, employee size sheet, artwork, notes", "Fields and uploads are present in the self-service and RFQ components."],
    ["RFQ and quote", "RFQ form with attachments, admin quote preparation, revisions, buyer approval or revision request", "Buyer dashboard and admin RFQ queue components are present."],
    ["Purchase order", "Buyer PO upload and admin/direct PO creation paths", "Purchase-order actions are present in buyer and finance/admin components."],
    ["Payments", "Razorpay advance, balance, requested payment, full upfront, payment confirmations, receipt vouchers", "Payment mutations, payment links, payment states, and receipt voucher routes are present."],
    ["Production and fulfilment", "QC gate, status updates, shipment workspace, dispatch and delivered states, replacement requests", "Order detail, QC, shipment, status timeline, and replacement components are present."],
    ["Finance documents", "Proforma, customer tax invoice, vendor PO, brand tax invoice, delivery challan, receipt voucher, settlement statement, order summary", "Templates and API routes are present for the listed documents."],
])
heading(doc, "2 Current corporate UI inventory")
add_table(doc, ["Screen or route","What it provides","Primary users"], [
    ["/corporate", "Corporate catalog, RFQ entry, corporate dashboard link, legacy self-service link", "Buyer"],
    ["/corporate-orders", "Legacy five-step order flow: Company, Product, Branding, Employee Sizes, Summary", "Buyer"],
    ["/profile/corporate", "Dashboard metrics for requests, quotes, orders, approved quotes, payments, deliveries; quote review and order setup", "Buyer"],
    ["/profile/corporate/request-quote", "Company/contact, requirements, quantity, budget, date, procurement mode, preferences, description, attachments", "Buyer"],
    ["/profile/corporate/self-service", "Self-service flow entry from the corporate dashboard", "Buyer"],
    ["/dashboard/general/corporate", "Corporate command center for RFQs, tasks, reports, finance, and executive views", "Renivet admin"],
    ["/dashboard/general/corporate-orders", "Corporate orders table, order detail, document chain, payment ledger, QC, shipment workspace", "Renivet operations/admin"],
    ["/dashboard/brands/{brandId}/corporate-orders", "Brand-side corporate orders workspace", "Supplier brand"],
])
para(doc, "Code evidence: the route and component inventory above is based on the current source tree under src/app, src/components/corporate-platform, src/components/corporate-orders, src/components/dashboard/general/corporate-orders, and src/components/profile/general.")
doc.add_page_break()

heading(doc, "3 End to end corporate order flow")
para(doc, "The current end-to-end flow can be represented as follows, from buyer entry through quote, order creation, payment, fulfilment, billing, and settlement.")
add_table(doc, ["Step","Buyer and system action","Current implementation","Current next state"], [
    ["1. Discover", "Buyer opens corporate surface and chooses catalog, RFQ, or self-service.", "Corporate landing page links all three paths.", "RFQ, catalog, or self-service entry."],
    ["2. Submit requirement", "Buyer enters company/contact, quantity, delivery date, use case, budget, preferences, description, and files.", "RFQ form captures these fields and accepts PDF, PPT, XLSX, images, and ZIP attachments within 50 MB.", "RFQ record and attachments are created."],
    ["3. Prepare quote", "Admin reviews RFQ, selects supplier/product inputs, sets cost/tax/advance, and creates a quote.", "Admin RFQ queue supports manual quote preparation and revisions.", "Quote is sent to buyer or revised."],
    ["4. Buyer decision", "Buyer approves quote, requests revision, or proceeds through PO/payment route.", "Buyer dashboard exposes quote decisions and purchase-order upload.", "Quote approved, revision requested, or PO/payment setup started."],
    ["5. Order release", "System creates order after accepted quote plus PO or required payment condition.", "Purchase-order and payment mutations are present; finance queue supports order release.", "Corporate order record and order snapshot are available to operations."],
    ["6. Advance and receipt", "Buyer pays advance or full amount; receipt voucher is generated.", "Razorpay flows, payment confirmation, balance links, and receipt voucher are implemented.", "Payment status and receipt voucher are recorded."],
    ["7. Supplier fulfilment", "Vendor PO, production, QC, shipment, dispatch, and delivery are tracked.", "Document chain, QC review, status timeline, and shipment workspace are present.", "Order progresses through production, QC, dispatch, and delivery statuses."],
    ["8. Final billing", "Customer tax invoice is issued; balance is collected where applicable.", "Tax invoice, balance payment link, payment status, and paid-in-full variant exist.", "Final tax invoice and balance status are recorded."],
    ["9. Settlement", "Supplier settlement is produced after delivery and finance checks.", "Settlement statement includes commission, GST on commission, TCS, TDS, and net remittance.", "Settlement statement and payout values are recorded."],
])
heading(doc, "4 Current features by capability")
for title, items in [
    ("Buyer and company data", ["Corporate profile and company snapshot", "Company name, contact person, email, phone, GSTIN, delivery address", "Employee count and employee size-sheet upload", "Customer notes and payment preference"]),
    ("Product and customization", ["Product type, GSM, fabric composition, colours", "Logo placement, print method, artwork upload", "Custom colours and extra charge rules", "Quantity and configured quote summary"]),
    ("Commercial workflow", ["RFQ creation and attachments", "Admin manual quote creation", "Quote revision history and buyer decision", "Purchase-order upload and direct admin order creation", "Partial advance or full upfront payment options"]),
    ("Operations", ["Corporate order status timeline", "Quality-control submission and approve/reject actions", "Ready for dispatch, dispatched, delivered, and replacement flows", "Shipment workspace and dispatch details", "Brand-side corporate order workspace"]),
    ("Finance and documents", ["Proforma invoice", "Customer tax invoice", "Vendor purchase order", "Receipt voucher", "Delivery challan", "Settlement statement", "Corporate order summary", "Payment ledger and balance payment links"]),
]:
    heading(doc, title, 2); bullets(doc, items)
doc.add_page_break()

heading(doc, "5 Current screens and document evidence")
para(doc, "The following captures are rendered from current repository previews and show the current document layouts and fields.")
heading(doc, "5.1 Application UI screenshots", 2)
para(doc, "The following screenshots were captured from the running local Renivet application using the current corporate test account. They show the current UI state and the five-step self-service order journey.")
ui_picture(doc, "01-corporate-entry.png", "Figure 1. Corporate entry page with catalog, Request Custom Quote, Open Corporate Dashboard, and Legacy Self-Service Flow entry points.")
ui_picture(doc, "02-self-service-step-1-company.png", "Figure 2. Self-service step 1 Company showing business information, contact, delivery, GSTIN, employee count, and shipping address fields.")
doc.add_page_break()
ui_picture(doc, "03-self-service-step-2-product.png", "Figure 3. Self-service step 2 Product showing the order progress panel and product selection stage.")
ui_picture(doc, "04-self-service-step-3-branding.png", "Figure 4. Self-service step 3 Branding showing logo placement and branding configuration stage.")
doc.add_page_break()
ui_picture(doc, "05-self-service-step-4-sizes.png", "Figure 5. Self-service step 4 Employee Sizes showing the employee size-list stage.")
ui_picture(doc, "06-self-service-step-5-summary.png", "Figure 6. Self-service step 5 Summary showing the final order review stage.")
doc.add_page_break()
ui_picture(doc, "07-rfq-request-form.png", "Figure 7. Request custom quote form showing Company and Contact Details, Order Requirements, Procurement Preferences, Requirement Description, Supporting Files, and next-step timeline.")
ui_picture(doc, "08-corporate-dashboard.png", "Figure 8. Corporate dashboard showing Self-Service Ordering, Request a Custom Quote, Company Profile, and Next Steps panels.")
doc.add_page_break()
heading(doc, "5.2 Admin UI screenshots", 2)
para(doc, "The following screenshots were captured from the current admin corporate-order workspaces using the same local development session.")
ui_picture(doc, "09-admin-corporate-hub.png", "Figure 9. Admin Corporate Hub showing Command Center, Requests (RFQs), Finance & POs, Task Control, Reports, Executive View, and current status counters.")
ui_picture(doc, "13-admin-corporate-orders-table.png", "Figure 10. Admin Corporate Orders table showing search, status filters, order type filter, quantities, order values, advance, balance, statuses, invoice download, and shipment actions.")
doc.add_page_break()
ui_picture(doc, "14-admin-order-detail-overview.png", "Figure 11. Admin order detail Overview showing corporate order status, payment status, order value, company and delivery data, supplier brand assignment, and commercial summary.")
ui_picture(doc, "15-admin-order-detail-documents.png", "Figure 12. Admin order detail Documents tab showing the order document workspace.")
doc.add_page_break()
ui_picture(doc, "16-admin-order-detail-payments.png", "Figure 13. Admin order detail Payments tab showing payment information and payment actions.")
ui_picture(doc, "17-admin-order-detail-activity.png", "Figure 14. Admin order detail Activity tab showing the order activity and status history workspace.")
doc.add_page_break()
heading(doc, "5.3 Corporate document screenshots", 2)
picture(doc, "corporate-order-summary.png", "Figure 15. Corporate order summary showing company information, garment configuration, branding/extras, order value, advance paid, and balance due.")
picture(doc, "corporate-proforma-invoice-preview.png", "Figure 16. Proforma invoice showing buyer/seller, quote reference, taxable value, GST, advance terms, and validity.")
doc.add_page_break()
picture(doc, "corporate-renivet-purchase-order-preview.png", "Figure 17. Renivet purchase order showing buyer, supplier, corporate order, customer PO, delivery mode, item, tax, and terms.")
picture(doc, "corporate-receipt-voucher-preview.png", "Figure 18. Receipt voucher showing advance received, payment reference, GST treatment, receipt total, and bank details.")
doc.add_page_break()
picture(doc, "corporate-delivery-challan-preview.png", "Figure 19. Delivery challan showing consignor, consignee, Renivet PO, e-way bill reference, goods movement, and terms.")
picture(doc, "corporate-settlement-statement-preview.png", "Figure 20. Settlement statement showing gross collection, embedded GST, platform commission, GST on commission, TCS, TDS, and net remittance.")
doc.add_page_break()
picture(doc, "corporate-full-payment-invoice-preview.png", "Figure 21. Customer tax invoice marked paid in full, including advance received, balance due, place of supply, GST breakup, QR, and declarations.")

heading(doc, "6 Document chain and ownership")
add_table(doc, ["Document","Purpose","Likely issuer","Trigger or prerequisite","Current evidence"], [
    ["Proforma invoice", "Commercial proposal before supply; not a tax invoice", "Renivet / corporate finance", "Approved quote or advance/order setup", "Rendered preview available"],
    ["Customer purchase order", "Buyer authorization and scope", "Corporate buyer", "Buyer chooses PO route", "Upload and admin creation paths present"],
    ["Receipt voucher", "Evidence of advance/payment received", "Renivet", "Payment confirmation", "Rendered preview available"],
    ["Vendor purchase order", "Instruction to supplier brand", "Renivet", "Order release and supplier details", "Rendered preview available"],
    ["Brand tax invoice", "Supplier invoice into Renivet process", "Supplier brand", "Supplier fulfilment and invoice upload", "Review and validation form present"],
    ["Delivery challan", "Goods movement document for direct fulfilment", "Renivet / operations", "Dispatch route requires goods-movement document", "Rendered preview available"],
    ["Customer tax invoice", "Final customer tax invoice", "Supplier brand or Renivet depending configured flow", "Supply/invoice issuance condition", "Paid-in-full and tax invoice previews available"],
    ["Settlement statement", "Supplier payout and deduction reconciliation", "Renivet finance", "Delivery/final billing and settlement checks", "Rendered preview available"],
])
doc.add_page_break()

heading(doc, "7 Current order states and statuses")
para(doc, "The corporate order record is used after the commercial setup is completed. The current UI exposes status, payment status, order value, advance paid, balance due, payment ledger, quality control, shipment, and document-chain information.")
add_table(doc, ["Current area","Statuses or values shown in the current UI"], [
    ["Commercial", "RFQ submitted, quote prepared, quote revision, quote approved, purchase order uploaded, order setup"],
    ["Payment", "Partial advance, full upfront, requested payment, paid, balance due, payment confirmation, receipt voucher"],
    ["Operations", "New, in production, QC pending, QC approved or rejected, ready for dispatch, dispatched, delivered"],
    ["Exceptions", "Replacement request, cancellation or status action areas where enabled by the order workspace"],
    ["Finance", "Proforma issued, supplier/brand invoice review, customer tax invoice, settlement statement, payout and deduction values"],
])

heading(doc, "8 Current system architecture")
para(doc, "The current corporate-order architecture is a Next.js application with server-rendered route entry points, client-side workflow components, tRPC procedures for interactive mutations and queries, service-layer business logic, database schemas and migrations, PDF document templates, and API download routes.")
add_table(doc, ["Layer","Current implementation evidence","Responsibility"], [
    ["Route layer", "src/app/(marketing)/corporate, corporate-orders, src/app/(protected)/profile/corporate, src/app/(protected)/dashboard/general/corporate, corporate-orders", "Page entry points, authentication gates, initial data loading, and workspace composition."],
    ["Buyer UI layer", "src/components/corporate-platform/customer-corporate-dashboard.tsx, request-quote-form.tsx, corporate-tabs.tsx; src/components/corporate-orders/corporate-order-page.tsx", "Buyer dashboard, RFQ form, self-service order wizard, quote decisions, PO upload, and payment actions."],
    ["Admin and brand UI layer", "admin-rfq-queue.tsx, admin-finance-queue.tsx, admin-task-center.tsx, admin-report-center.tsx, brand-corporate-workspace.tsx", "RFQ preparation, quote revisions, finance actions, operational tasks, reports, and supplier-brand workspaces."],
    ["Order operations UI", "corporate-order-detail.tsx and corporate-document-chain-panel.tsx", "Order snapshot, status changes, payment ledger, QC, shipment, file access, and document issuance."],
    ["tRPC/API layer", "src/lib/trpc/routes/general/corporate-platform.ts, corporate-orders.ts; src/app/api/corporate-orders/[id]/*", "Queries and mutations for corporate profiles, RFQs, quotes, orders, payments, statuses, and document downloads."],
    ["Service layer", "src/lib/services/corporate-platform.ts, corporate-order.ts, corporate-documents.ts, corporate-payment-request.tsx", "Corporate workflow orchestration, order operations, document numbering/rendering, and payment-request handling."],
    ["Validation and finance layer", "src/lib/validations/corporate-order.ts, corporate-platform.ts; src/lib/finance/corporate-place-of-supply.ts, corporate-tax-classification.ts", "Input validation, corporate quote/order validation, tax classification, place-of-supply handling, and tax calculations."],
    ["Data layer", "src/lib/db/schema/corporate-order.ts, corporate-platform.ts; drizzle migrations", "Corporate profiles, RFQs, quotes, purchase orders, payments, invoices, receipts, fulfilment, QC, settlement, and status history."],
    ["Document layer", "src/components/pdf/corporate-commercial-document-template.tsx, corporate-tax-invoice-template.tsx, corporate-settlement-statement-template.tsx, corporate-receipt-voucher-template.tsx, corporate-order-summary-template.tsx", "Renderable document layouts used by corporate PDF routes."],
])
heading(doc, "9 Current data and control flow")
add_table(doc, ["Flow","Current path"], [
    ["Buyer RFQ", "Buyer route -> request-quote-form -> corporate platform tRPC mutation -> corporate profile/RFQ/attachment records -> admin RFQ queue."],
    ["Admin quote", "Admin RFQ queue -> quote preparation or manual quote modal -> quote record and revision history -> buyer corporate dashboard."],
    ["Buyer approval", "Buyer dashboard -> quote decision -> accepted quote state -> purchase-order upload or payment setup."],
    ["Order creation", "Purchase order or payment mutation -> corporate order service -> corporate order plus company, delivery, product, customization, and commercial snapshots."],
    ["Payment", "Buyer payment action -> Razorpay checkout -> payment confirmation mutation -> corporate payment record -> receipt voucher and updated order payment state."],
    ["Fulfilment", "Order detail -> vendor PO/document chain -> supplier/brand workspace -> QC review -> status timeline -> shipment and dispatch workspace."],
    ["Final billing", "Order/document chain -> customer tax invoice route/template -> invoice number and tax values -> buyer/admin download."],
    ["Settlement", "Delivered/finalized order data -> settlement calculation -> settlement statement route/template -> finance and supplier records."],
])
heading(doc, "10 Evidence-backed implementation inventory")
add_table(doc, ["Repository area","Evidence found"], [
    ["Buyer corporate platform", "src/components/corporate-platform/customer-corporate-dashboard.tsx, request-quote-form.tsx, corporate-tabs.tsx"],
    ["Legacy order journey", "src/components/corporate-orders/corporate-order-page.tsx with five-step flow and uploads"],
    ["Admin platform", "admin-rfq-queue.tsx, admin-finance-queue.tsx, admin-task-center.tsx, admin-report-center.tsx"],
    ["Order operations", "corporate-order-detail.tsx and corporate-document-chain-panel.tsx"],
    ["Document templates", "src/components/pdf/corporate-commercial-document-template.tsx, corporate-tax-invoice-template.tsx, corporate-settlement-statement-template.tsx, corporate-receipt-voucher-template.tsx, corporate-order-summary-template.tsx"],
    ["API document routes", "src/app/api/corporate-orders/[id]/... routes for invoice, summary, receipt, vendor PO, challan, settlement, and fulfilment order"],
    ["Rendered previews", "output/pdf/*.pdf and tmp/corp-docs-png/*.png used in this report"],
])
heading(doc, "Appendix A Evidence note")
para(doc, "This document records the implementation visible in the repository and the rendered corporate document previews available at review time. The invoice and document previews contain test or sample identities and values and are shown to document the current layouts, fields, and document chain.")
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
